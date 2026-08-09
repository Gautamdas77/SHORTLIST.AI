import json
import os
from io import BytesIO

from docx import Document
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field, field_validator
from pypdf import PdfReader

load_dotenv()
API_KEY = os.getenv("GROQ_API_KEY")
MODEL = "llama-3.3-70b-versatile"


class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]


class MatchResult(BaseModel):
    score: float
    details: dict


class Experience(BaseModel):
    company: str | None = None
    position: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)

    @field_validator("skills_used", mode="before")
    @classmethod
    def empty_skills_for_null(cls, value):
        return [] if value is None else value


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience: float | None = None
    education: list[dict] = Field(default_factory=list)
    experience: list[Experience] = Field(default_factory=list)
    skills: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)


def get_client():
    if not API_KEY:
        raise ValueError("GROQ_API_KEY is not configured")
    return Groq(api_key=API_KEY)


def parse_resume(resume_text: str) -> Resume:
    schema = Resume.model_json_schema()
    messages = [
        {"role": "system", "content": f"You are an expert resume parser. Extract information based on meaning, including internships, projects, and skills mentioned across the entire resume. Return ONLY valid JSON matching this schema: {schema}. Do not invent information. Missing values are null and missing lists are empty."},
        {"role": "user", "content": f"Parse this resume:\n{resume_text}"},
    ]
    response = get_client().chat.completions.create(model=MODEL, messages=messages, response_format={"type": "json_object"})
    return Resume(**json.loads(response.choices[0].message.content))


def final_score(job: JobD, resume: Resume) -> MatchResult:
    schema = MatchResult.model_json_schema()
    prompt = f"""Compare this candidate with the job description and return JSON matching this schema: {schema}
Include matching skills, missing important skills, whether experience and education requirements are met, a 0-100 overall match percentage, and a short final verdict.
JOB DESCRIPTION:\n{job.model_dump_json(indent=2)}\nCANDIDATE RESUME:\n{resume.model_dump_json(indent=2)}"""
    response = get_client().chat.completions.create(model=MODEL, messages=[{"role": "user", "content": prompt}], response_format={"type": "json_object"})
    return MatchResult(**json.loads(response.choices[0].message.content))


def read_pdf(file_bytes: bytes) -> str:
    text = []
    for page in PdfReader(BytesIO(file_bytes)).pages:
        if page_text := page.extract_text():
            text.append(page_text)
    return "\n".join(text)


def read_docx(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    text = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    text.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip())
    return "\n".join(text)


def read_resume(file_name: str, file_bytes: bytes) -> str:
    suffix = os.path.splitext(file_name)[1].lower()
    if suffix == ".pdf":
        return read_pdf(file_bytes)
    if suffix == ".docx":
        return read_docx(file_bytes)
    if suffix == ".txt":
        return file_bytes.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {suffix}")


def analyze_candidates(job_description: str, files: list[tuple[str, bytes]]) -> list[dict]:
    if not job_description.strip():
        raise ValueError("A job description is required")
    if not files:
        raise ValueError("Upload at least one resume")

    schema = JobD.model_json_schema()
    messages = [
        {"role": "system", "content": f"You are an expert HR assistant. Extract structured information from the job description. Return ONLY valid JSON matching this schema: {schema}. Do not invent information."},
        {"role": "user", "content": f"Analyze this job description:\n{job_description}"},
    ]
    job_response = get_client().chat.completions.create(model=MODEL, messages=messages, response_format={"type": "json_object"})
    job = JobD(**json.loads(job_response.choices[0].message.content))

    results = []
    for file_name, file_bytes in files:
        parsed = parse_resume(read_resume(file_name, file_bytes))
        scored = final_score(job, parsed)
        results.append({
            "name": parsed.name or os.path.splitext(file_name)[0],
            "file_name": file_name,
            "score": max(0, min(100, round(scored.score))),
            "details": scored.details,
            "skills": parsed.skills,
            "experience": parsed.total_experience,
        })
    return sorted(results, key=lambda result: result["score"], reverse=True)
